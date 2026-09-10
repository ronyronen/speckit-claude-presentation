#!/usr/bin/env python3
"""Builds presentation/speckit-workshop-draft-en.pptx -- an all-English
version of the deck. Diagrams come from presentation/assets/generated_en/
(pure LTR, no bidi needed); terminal/code content cards come from
presentation/assets/generated/ since they were already English-only
content (see make_content_cards.py) regardless of which language deck
uses them. Run from the repository root:

    python3 presentation/build_deck_en.py

Hardware-dependent slides are explicitly marked PENDING and contain no
fabricated data -- they will be completed after checkpoint
09-hardware-verified.
"""
import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

ASSETS = "presentation/assets/generated"
ASSETS_EN = "presentation/assets/generated_en"
OUT = "presentation/speckit-workshop-draft-en.pptx"

NAVY = RGBColor(0x0F, 0x17, 0x2A)
BLUE = RGBColor(0x25, 0x63, 0xEB)
SLATE = RGBColor(0x33, 0x41, 0x55)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
AMBER = RGBColor(0xD9, 0x77, 0x06)
RED = RGBColor(0xB9, 0x1C, 0x1C)
GREEN = RGBColor(0x15, 0x80, 0x3D)
LIGHT_BG = RGBColor(0xF8, 0xFA, 0xFC)
SECTION_BG = RGBColor(0x1E, 0x1B, 0x4B)

FONT = "Arial"

SW, SH = Inches(13.333), Inches(7.5)


def add_slide(prs, bg=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    if bg is not None:
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = bg
    return slide


def add_title(slide, text, top=Inches(0.35), color=NAVY, size=30, height=Inches(0.9)):
    box = slide.shapes.add_textbox(Inches(0.5), top, SW - Inches(1.0), height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.name = FONT
    r.font.color.rgb = color
    return box


def add_bullets(slide, items, top=Inches(1.4), left=Inches(0.6), width=None, height=Inches(4.6), size=20):
    width = width or (SW - Inches(1.2))
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(12)
        text, kwargs = (item, {}) if isinstance(item, str) else (item["text"], item)
        r = p.add_run()
        r.text = text
        r.font.size = Pt(kwargs.get("size", size))
        r.font.bold = kwargs.get("bold", False)
        r.font.italic = kwargs.get("italic", False)
        r.font.name = FONT
        r.font.color.rgb = kwargs.get("color", SLATE)
    return box


def add_note(slide, text):
    notes = slide.notes_slide
    notes.notes_text_frame.text = text


def add_pending_banner(slide, text="Pending physical hardware verification (checkpoint 09-hardware-verified)"):
    box = slide.shapes.add_textbox(Inches(0.6), SH - Inches(0.85), SW - Inches(1.2), Inches(0.55))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "⏳ " + text
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.italic = True
    r.font.name = FONT
    r.font.color.rgb = AMBER
    rect = slide.shapes.add_shape(1, Inches(0.5), SH - Inches(0.95), SW - Inches(1.0), Inches(0.7))
    rect.fill.solid()
    rect.fill.fore_color.rgb = RGBColor(0xFE, 0xF3, 0xC7)
    rect.line.color.rgb = AMBER
    rect.line.width = Pt(1.5)
    rect.shadow.inherit = False
    spTree = slide.shapes._spTree
    spTree.remove(rect._element)
    spTree.insert(2, rect._element)
    return box


def add_missing_visual_placeholder(slide, label, left, top, width, height):
    rect = slide.shapes.add_shape(1, left, top, width, height)
    rect.fill.solid()
    rect.fill.fore_color.rgb = RGBColor(0xF1, 0xF5, 0xF9)
    rect.line.color.rgb = RGBColor(0x94, 0xA3, 0xB8)
    rect.line.width = Pt(1.5)
    rect.line.dash_style = 4
    tf = rect.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "\U0001F4F7 " + label
    r.font.size = Pt(16)
    r.font.italic = True
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run()
    r2.text = "(real screenshot missing -- see screenshot-checklist.md)"
    r2.font.size = Pt(13)
    r2.font.italic = True
    r2.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
    return rect


def add_image(slide, path, left, top, width=None, height=None, en=False):
    base = ASSETS_EN if en else ASSETS
    full = os.path.join(base, path)
    return slide.shapes.add_picture(full, left, top, width=width, height=height)


def main():
    prs = Presentation()
    prs.slide_width = SW
    prs.slide_height = SH

    # ---------- Slide 1: Title ----------
    s = add_slide(prs, bg=NAVY)
    box = s.shapes.add_textbox(Inches(1), Inches(2.3), SW - Inches(2), Inches(1.6))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "From Prompt to Product"
    r.font.size = Pt(48)
    r.font.bold = True
    r.font.color.rgb = WHITE
    r.font.name = FONT
    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run()
    r2.text = "Spec-Driven Development with GitHub Spec Kit"
    r2.font.size = Pt(28)
    r2.font.color.rgb = RGBColor(0xA5, 0xB4, 0xFC)
    r2.font.name = FONT

    box2 = s.shapes.add_textbox(Inches(1), Inches(4.3), SW - Inches(2), Inches(0.8))
    tf2 = box2.text_frame
    p3 = tf2.paragraphs[0]
    p3.alignment = PP_ALIGN.CENTER
    r3 = p3.add_run()
    r3.text = "A real example on the Heltec WiFi Kit V3"
    r3.font.size = Pt(20)
    r3.font.italic = True
    r3.font.color.rgb = RGBColor(0xCB, 0xD5, 0xE1)
    r3.font.name = FONT
    add_note(s, "Open with a question to the audience: who has ever asked an AI to write code straight from a prompt?")

    # ---------- Slide 2: The vague prompt ----------
    s = add_slide(prs)
    add_title(s, "The Vague Prompt")
    rect = s.shapes.add_shape(1, Inches(1.0), Inches(1.7), SW - Inches(2.0), Inches(1.8))
    rect.fill.solid()
    rect.fill.fore_color.rgb = LIGHT_BG
    rect.line.color.rgb = RGBColor(0x94, 0xA3, 0xB8)
    box = s.shapes.add_textbox(Inches(1.2), Inches(1.8), SW - Inches(2.4), Inches(2))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = ("Build an application for a Heltec WiFi Kit V3 that reads a sensor,\n"
              "shows the value on the display, and shows a warning when the value is too high.")
    r.font.size = Pt(22)
    r.font.name = "Courier New"
    r.font.color.rgb = SLATE
    box3 = s.shapes.add_textbox(Inches(1.2), Inches(4.2), SW - Inches(2.4), Inches(1))
    tf3 = box3.text_frame
    tf3.word_wrap = True
    p3 = tf3.paragraphs[0]
    p3.alignment = PP_ALIGN.CENTER
    r3 = p3.add_run()
    r3.text = "❓ Is this enough information to build the correct product?"
    r3.font.size = Pt(28)
    r3.font.bold = True
    r3.font.color.rgb = RED
    r3.font.name = FONT
    add_note(s, "Open question to the audience -- give them a moment to think before continuing.")

    # ---------- Slide 3: Open questions ----------
    s = add_slide(prs)
    add_title(s, "Everything the Prompt Never Answered")
    add_bullets(s, [
        "Which sensor?",
        "How often is it sampled?",
        "What exactly is \"too high\"?",
        "What happens exactly at the threshold?",
        "When does the warning clear?",
        "What happens if the sensor fails?",
        "What does the display show during an error?",
        "What gets logged to Serial?",
    ], size=22)
    add_note(s, "Can briefly show examples/prompt-only/naive_monitor.ino here.")

    # ---------- Slide 4: Core message ----------
    s = add_slide(prs, bg=NAVY)
    box = s.shapes.add_textbox(Inches(1), Inches(2.6), SW - Inches(2), Inches(2.2))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "The problem is usually not that the AI cannot write the code."
    r.font.size = Pt(30)
    r.font.bold = True
    r.font.color.rgb = WHITE
    r.font.name = FONT
    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run()
    r2.text = "The problem is that the product behavior was not sufficiently defined."
    r2.font.size = Pt(30)
    r2.font.bold = True
    r2.font.color.rgb = RGBColor(0xFC, 0xD3, 0x4D)
    r2.font.name = FONT
    add_note(s, "This is the one sentence the audience should remember from the whole workshop. Pause here for a beat.")

    # ---------- Slide 5: Prompt loop ----------
    s = add_slide(prs)
    add_title(s, "From Prompt to an Infinite Loop")
    add_image(s, "diagram-prompt-vs-spec-en.png", Inches(0.5), Inches(1.3), width=SW - Inches(1.0), en=True)
    add_note(s, "Focus on the top half of the diagram for now. Every 'Fix' is a product decision made after code already exists that contradicts it.")

    # ---------- Slide 6: Spec-driven alternative ----------
    s = add_slide(prs)
    add_title(s, "The Alternative: A Spec-Driven Process")
    add_image(s, "diagram-prompt-vs-spec-en.png", Inches(0.5), Inches(1.3), width=SW - Inches(1.0), en=True)
    add_note(s, "Now focus on the bottom half. The difference isn't 'more bureaucracy' -- it's moving decisions to where they're cheap to make.")

    # ---------- Slide 7: What each stage catches ----------
    s = add_slide(prs)
    add_title(s, "What Each Stage Catches")
    rows = [
        ("Specification", "We never agreed what this does"),
        ("Clarify", "We agreed on the happy path, not the edges"),
        ("Plan", "We agreed on what, not how"),
        ("Analyze", "Two of our own documents contradict each other"),
    ]
    top = Inches(1.6)
    for term, desc in rows:
        box = s.shapes.add_textbox(Inches(0.8), top, SW - Inches(1.6), Inches(0.8))
        tf = box.text_frame
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        r = p.add_run()
        r.text = term
        r.font.size = Pt(24)
        r.font.bold = True
        r.font.name = "Courier New"
        r.font.color.rgb = BLUE
        r2 = p.add_run()
        r2.text = f"    →    {desc}"
        r2.font.size = Pt(24)
        r2.font.name = FONT
        r2.font.color.rgb = SLATE
        top += Inches(1.0)

    # ---------- Slide 8: Spec Kit definition ----------
    s = add_slide(prs)
    add_title(s, "What Is GitHub Spec Kit")
    add_bullets(s, [
        "Spec Kit = a CLI tool (specify) + a set of Skills/Commands for a coding agent",
        "It doesn't write code itself -- it structures the *process* that leads to writing the code",
    ], size=24)
    add_missing_visual_placeholder(s, "Spec Kit website + GitHub repo", Inches(0.8), Inches(3.2), SW - Inches(1.6), Inches(2.8))
    add_note(s, "screenshots #1-2 from the checklist.")

    # ---------- Slide 9: Installation ----------
    s = add_slide(prs)
    add_title(s, "Installation")
    box = s.shapes.add_textbox(Inches(1.2), Inches(1.5), SW - Inches(2.4), Inches(1.0))
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = "uv tool install specify-cli\nspecify init --here --integration claude"
    r.font.size = Pt(20)
    r.font.name = "Courier New"
    r.font.color.rgb = SLATE
    add_image(s, "terminal-specify-init.png", Inches(1.0), Inches(2.7), width=Inches(11.3))
    add_note(s, "This is the real output from this repo's own initialization.")

    # ---------- Slide 10: 7-stage workflow ----------
    s = add_slide(prs)
    add_title(s, "The Workflow in Seven Stages")
    add_image(s, "diagram-speckit-workflow-en.png", Inches(0.6), Inches(1.6), width=SW - Inches(1.2), en=True)
    add_note(s, "Don't dive into command syntax -- the goal is the mental model.")

    # ---------- Slide 11: Real files ----------
    s = add_slide(prs)
    add_title(s, "The Output: Real Files in the Repository")
    add_bullets(s, [
        {"text": "specs/001-heltec-monitor/{spec,plan,tasks}.md", "size": 24, "color": BLUE, "bold": True},
        "Not a chat conversation that disappears",
        "This is what lets you switch agents, come back a month later, and see exactly why each decision was made",
    ], size=22)
    add_note(s, "")

    # ---------- Slide 12: who does what ----------
    s = add_slide(prs)
    add_title(s, "Who Does What")
    add_image(s, "diagram-tool-relationship-en.png", Inches(2.0), Inches(1.2), height=Inches(6.0), en=True)
    add_note(s, "'Spec Kit is not the coding agent. Claude Code or Codex do the work. The IDE is the human workspace.'")

    # ---------- Slide 13: same repo two agents ----------
    s = add_slide(prs)
    add_title(s, "Same Repo, Two Agents")
    add_bullets(s, [
        {"text": "Claude Code  →  /speckit-specify   (.claude/skills/)", "size": 22, "color": BLUE},
        {"text": "Codex  →  $speckit-specify   (.agents/skills/)", "size": 22, "color": AMBER},
        "This isn't theoretical -- this repo has both integrations installed side by side",
    ], top=Inches(1.6))
    add_missing_visual_placeholder(s, "Claude Code + Codex session (same repo)", Inches(0.8), Inches(4.0), SW - Inches(1.6), Inches(2.4))
    add_note(s, "screenshots #6-7.")

    # ---------- Slide 14: VS Code briefly ----------
    s = add_slide(prs)
    add_title(s, "VS Code, Briefly")
    add_bullets(s, [
        "Extensions: PlatformIO IDE, C/C++ -- each with a specific reason, not a long list",
        "JetBrains/CLion: the same files, a different terminal",
    ], size=24)
    add_note(s, "")

    # ---------- Slide 15: comparison table ----------
    s = add_slide(prs)
    add_title(s, "Direct Prompting vs. Spec Kit")
    rows = [
        ("Requirements", "Live in chat history", "Live as repository files"),
        ("Clarifications", "Agent assumes on its own", "Explicit clarification stage"),
        ("Traceability", "Low", "High"),
        ("Repeatability", "Low-to-medium", "High"),
    ]
    top = Inches(1.7)
    for topic, direct, spec in rows:
        box = s.shapes.add_textbox(Inches(0.8), top, SW - Inches(1.6), Inches(0.7))
        tf = box.text_frame
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        r = p.add_run()
        r.text = topic
        r.font.bold = True
        r.font.size = Pt(20)
        r.font.name = "Courier New"
        r.font.color.rgb = BLUE
        r2 = p.add_run()
        r2.text = f"   —   {direct}   →   {spec}"
        r2.font.size = Pt(20)
        r2.font.name = FONT
        r2.font.color.rgb = SLATE
        top += Inches(0.85)
    add_note(s, "")

    # ---------- Slide 16: alternatives ----------
    s = add_slide(prs)
    add_title(s, 'And Why Not Just "Another Framework"?')
    add_bullets(s, [
        "OpenSpec -- similar territory, a different template/command set",
        "BMAD -- more roles and handoffs, more overhead",
        "Spec Kit sits between free-form prompting and heavy multi-agent processes",
    ], size=22)
    add_note(s, "Not a framework survey -- that's not the point of this workshop.")

    # ---------- Slide 17: hardware verified ----------
    s = add_slide(prs)
    add_title(s, "The Hardware, Verified")
    add_bullets(s, [
        "Heltec WiFi Kit V3 (ESP32-S3), built-in OLED, external BME280",
        {"text": "No built-in environmental sensor on the board -- that's already a product decision", "bold": True, "color": RED},
        "Surprising technical detail: the display's I2C pins are 17/18, not the ESP32 default (21/22)",
    ], size=22)
    add_note(s, "")

    # ---------- Slide 18: Stage A ----------
    s = add_slide(prs)
    add_title(s, "Stage A: Prompt Only")
    add_bullets(s, [
        {"text": "The point: implicit assumption vs. explicit decision", "bold": True, "size": 24, "color": BLUE},
        "Hardware wiring (I2C, Vext) is correct in both versions",
        "What's missing: which sensor, the exact threshold, whether hysteresis is needed",
        {"text": "Key example: no hysteresis -- not a \"bug,\" the result of a requirement that was never stated", "italic": True},
    ], size=20)
    add_note(s, "Visual: examples/prompt-only/naive_monitor.ino")

    # ---------- Slide 19: Specify + Clarify ----------
    s = add_slide(prs)
    add_title(s, "Stage B-C: Specify + Clarify")
    add_bullets(s, [
        "The 4 questions actually asked:",
        "Sampling rate • exact threshold bounds • failure debounce • error display content",
    ], size=22)
    add_image(s, "code-spec-clarifications.png", Inches(0.9), Inches(3.0), width=Inches(11.5))
    add_note(s, "")

    # ---------- Slide 20: Hysteresis ----------
    s = add_slide(prs)
    add_title(s, "Hysteresis: The Clearest Example")
    add_image(s, "diagram-hysteresis-en.png", Inches(0.6), Inches(1.3), width=SW - Inches(1.2), en=True)
    add_note(s, "This is not a programming problem. It's a product decision that had to be made *before* writing code.")

    # ---------- Slide 21: Plan + Tasks ----------
    s = add_slide(prs)
    add_title(s, "Stage D-E: Plan + Tasks")
    add_image(s, "diagram-architecture-en.png", Inches(0.6), Inches(1.5), width=SW - Inches(1.2), en=True)
    add_note(s, "screenshots #10-11 (plan.md, tasks.md) if you want to also show the files themselves.")

    # ---------- Slide 22: Analyze catch ----------
    s = add_slide(prs)
    add_title(s, "⭐ Stage F: Analyze Finds a Real Contradiction", size=28)
    add_image(s, "diagram-analyze-catch-en.png", Inches(2.6), Inches(1.15), height=Inches(6.1), en=True)
    add_note(s, ("Most important slide -- don't rush through it. This is not a made-up example for the talk -- "
                 "it really happened while building this project. Can show live: git show b91cd94 -- specs/001-heltec-monitor/spec.md"))

    # ---------- Slide 22b: Analyze message ----------
    s = add_slide(prs, bg=NAVY)
    box = s.shapes.add_textbox(Inches(1), Inches(2.2), SW - Inches(2), Inches(2.6))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "/speckit-analyze didn't find a syntax problem or a code bug."
    r.font.size = Pt(28)
    r.font.bold = True
    r.font.color.rgb = WHITE
    r.font.name = FONT
    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run()
    r2.text = "It found that our own requirements contradicted each other -- before a single line of code was written."
    r2.font.size = Pt(26)
    r2.font.bold = True
    r2.font.color.rgb = RGBColor(0xFC, 0xD3, 0x4D)
    r2.font.name = FONT
    add_image(s, "code-analyze-diff.png", Inches(1.5), Inches(4.5), width=Inches(10.3))
    add_note(s, "The strongest demonstration of Spec Kit's value in the entire talk.")

    # ---------- Slide 23: Implement ----------
    s = add_slide(prs)
    add_title(s, "Stage G: Implement")
    add_image(s, "code-monitor-logic.png", Inches(0.9), Inches(1.25), width=Inches(10.8))
    add_image(s, "terminal-pio-test.png", Inches(0.9), Inches(5.15), width=Inches(4.6))
    box = s.shapes.add_textbox(Inches(5.8), Inches(5.35), Inches(6.6), Inches(1.8))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = "monitor_logic is fully tested with zero hardware"
    r.font.size = Pt(20)
    r.font.bold = True
    r.font.name = FONT
    r.font.color.rgb = GREEN
    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.LEFT
    r2 = p2.add_run()
    r2.text = "(pio test -e native, 11/11 passed)"
    r2.font.size = Pt(18)
    r2.font.name = "Courier New"
    r2.font.color.rgb = GREEN
    add_note(s, "")

    # ---------- Slide 24: Live demo (PENDING) ----------
    s = add_slide(prs)
    add_title(s, "Live Demo")
    add_bullets(s, [
        "Normal temperature → OK",
        "Warming above 30°C → WARNING",
        {"text": "Partial cooling (28-30°C) → stays WARNING (the single most important point)", "bold": True, "color": RED},
        "Cooling below 28°C → back to OK",
        "If there's time: disconnect the sensor → SENSOR_ERROR",
    ], top=Inches(1.5), size=20)
    add_missing_visual_placeholder(s, "Serial output + OLED photo", Inches(0.8), Inches(4.6), SW - Inches(1.6), Inches(1.7))
    add_pending_banner(s)
    add_note(s, "Everything on this slide depends on the physical hardware test, which hasn't happened yet.")

    # ---------- Slide 25: Git checkpoints ----------
    s = add_slide(prs)
    add_title(s, "Git Checkpoints")
    add_image(s, "terminal-git-log.png", Inches(0.9), Inches(1.4), width=Inches(11.5))
    add_pending_banner(s, "Tag 09-hardware-verified does not exist yet -- pending physical test")
    add_note(s, "Every stage already exists as a tag in the repository -- the demo doesn't depend on a live AI run that could fail on stage.")

    # ---------- Slide 26: Key takeaways ----------
    s = add_slide(prs, bg=NAVY)
    add_title(s, "Five Things to Remember", color=WHITE, size=32)
    add_bullets(s, [
        {"text": "•  AI coding speed does not remove the need for clear requirements", "color": WHITE, "size": 22},
        {"text": "•  Spec Kit gives a repeatable path from idea to implementation", "color": WHITE, "size": 22},
        {"text": "•  Claude Code and Codex are the agents; Spec Kit is the process and structure", "color": WHITE, "size": 22},
        {"text": "•  The spec, plan, and tasks live in the repository -- they don't disappear into a chat", "color": WHITE, "size": 22},
        {"text": "•  The best proof: not that it generates code, but that it forces answers to important questions before the wrong code gets written", "color": RGBColor(0xFC, 0xD3, 0x4D), "size": 22, "bold": True},
    ], top=Inches(1.6))
    add_note(s, "")

    # ---------- Slide 27: Links ----------
    s = add_slide(prs)
    add_title(s, "Links and Credits")
    add_bullets(s, [
        "Spec Kit: github.com/github/spec-kit",
        "Spec Kit docs: github.github.io/spec-kit",
        "Claude Code docs: docs.claude.com/en/docs/claude-code",
        "Codex docs: developers.openai.com/codex",
        "PlatformIO: platformio.org",
        "Heltec WiFi Kit 32 (V3): heltec.org/project/wifi-kit32-v3",
        {"text": "Companion repository: (to be filled in after publishing)", "italic": True},
    ], size=20)
    add_note(s, "See presentation/slide-links.md for the full list.")

    prs.save(OUT)
    print(f"saved {OUT} with {len(prs.slides.__iter__.__self__._sldIdLst)} slides")


if __name__ == "__main__":
    main()
